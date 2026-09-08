from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document


def export_requirements_word(requirements: list[dict[str, Any]], document: dict[str, Any], output_path: Path | str) -> Path:
    """Exporta documentos y requisitos a Word con evidencia y página de origen."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("Documentos y requisitos - SECOP Analyzer", level=1)
    doc.add_paragraph(f"Total de páginas: {document.get('total_paginas', 'N/A')}")
    doc.add_paragraph(f"Página inicio de ítems: {document.get('item_start_page', 'N/A')}")
    doc.add_paragraph(f"Lugar principal de ejecución: {document.get('lugar_principal_ejecucion', 'Taller físico o sitio de trabajo')}")

    for idx, item in enumerate(requirements, start=1):
        doc.add_heading(f"{idx}. {item.get('nombre', 'No encontrado')}", level=2)
        doc.add_paragraph(f"Acción: {item.get('accion', 'INFORMACIÓN')}")
        doc.add_paragraph(f"Obligatorio: {'Sí' if item.get('obligatorio') else 'No'}")
        doc.add_paragraph(f"Página: {item.get('pagina', 'N/A')}")
        doc.add_paragraph(f"Evidencia: {item.get('evidencia', 'No encontrado')}")
        doc.add_paragraph(f"Confianza: {item.get('confianza', 'N/A')}")

    doc.save(output)
    return output
